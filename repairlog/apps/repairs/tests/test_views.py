from datetime import timedelta

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from apps.clients.models import Client
from apps.equipment.models import Brand, Equipment, EquipmentType
from apps.repairs.models import Repair

User = get_user_model()


class RepairViewsTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="master", password="pass12345")
        self.client.force_login(self.user)
        self.client_obj = Client.objects.create(name="Ресторан Восток", phone="+79990001122")
        self.equipment_type = EquipmentType.objects.create(
            name="Пароконвектомат", category=EquipmentType.Category.KITCHEN
        )
        self.brand = Brand.objects.create(name="Rational")
        self.equipment = Equipment.objects.create(
            client=self.client_obj,
            equipment_type=self.equipment_type,
            brand=self.brand,
            model="SCC 61",
            serial_number="SN-001",
        )


class RepairWizardClientStepViewTests(RepairViewsTestCase):
    def test_requires_login(self):
        self.client.logout()
        response = self.client.get(reverse("repairs:wizard_client"))
        self.assertEqual(response.status_code, 302)

    def test_lists_clients(self):
        response = self.client.get(reverse("repairs:wizard_client"))
        self.assertContains(response, "Ресторан Восток")

    def test_search_by_name(self):
        Client.objects.create(name="Кафе Плюс")
        response = self.client.get(reverse("repairs:wizard_client"), {"q": "Восток"})
        self.assertContains(response, "Ресторан Восток")
        self.assertNotContains(response, "Кафе Плюс")


class RepairWizardEquipmentStepViewTests(RepairViewsTestCase):
    def test_shows_client_equipment(self):
        response = self.client.get(
            reverse("repairs:wizard_equipment", args=[self.client_obj.pk])
        )
        self.assertContains(response, "SCC 61")

    def test_add_equipment_url_routes_back_into_wizard(self):
        response = self.client.get(
            reverse("repairs:wizard_equipment", args=[self.client_obj.pk])
        )
        add_url = response.context["add_equipment_url"]
        self.assertIn(reverse("equipment:create"), add_url)
        self.assertIn(f"client={self.client_obj.pk}", add_url)
        self.assertIn("next=", add_url)


class RepairCreateViewTests(RepairViewsTestCase):
    def test_requires_equipment_param(self):
        response = self.client.get(reverse("repairs:create"))
        self.assertEqual(response.status_code, 404)

    def test_prefills_master_and_reported_at(self):
        response = self.client.get(
            reverse("repairs:create"), {"equipment": self.equipment.pk}
        )
        self.assertEqual(response.context["form"].initial["master"], self.user.pk)
        self.assertEqual(response.context["form"].initial["reported_at"], timezone.localdate())

    def test_creates_repair_with_new_status(self):
        response = self.client.post(
            f"{reverse('repairs:create')}?equipment={self.equipment.pk}",
            {
                "equipment": self.equipment.pk,
                "master": self.user.pk,
                "reported_at": timezone.localdate().isoformat(),
                "symptom": "Не набирает температуру",
                "error_code": "E4",
            },
        )
        repair = Repair.objects.get(symptom="Не набирает температуру")
        self.assertRedirects(response, reverse("repairs:detail", args=[repair.pk]))
        self.assertEqual(repair.equipment, self.equipment)
        self.assertEqual(repair.status, Repair.Status.NEW)

    def test_error_code_field_wired_for_fault_suggestions(self):
        response = self.client.get(
            reverse("repairs:create"), {"equipment": self.equipment.pk}
        )
        form = response.context["form"]
        attrs = form.fields["error_code"].widget.attrs
        self.assertEqual(attrs["hx-get"], reverse("knowledge:fault_suggestions"))
        self.assertIn(str(self.equipment_type.pk), attrs["hx-vals"])
        self.assertContains(response, 'id="fault-suggestions"')


class RepairListViewTests(RepairViewsTestCase):
    def setUp(self):
        super().setUp()
        self.repair = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Не набирает температуру",
            error_code="E4",
        )

    def test_requires_login(self):
        self.client.logout()
        response = self.client.get(reverse("repairs:list"))
        self.assertEqual(response.status_code, 302)

    def test_lists_active_repairs(self):
        response = self.client.get(reverse("repairs:list"))
        self.assertContains(response, "Не набирает температуру")

    def test_archived_hidden_by_default(self):
        self.repair.is_archived = True
        self.repair.save(update_fields=["is_archived"])
        response = self.client.get(reverse("repairs:list"))
        self.assertNotContains(response, "Не набирает температуру")

    def test_archived_shown_with_filter(self):
        self.repair.is_archived = True
        self.repair.save(update_fields=["is_archived"])
        response = self.client.get(reverse("repairs:list"), {"archived": "1"})
        self.assertContains(response, "Не набирает температуру")

    def test_filter_by_status(self):
        other = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Другая заявка",
            status=Repair.Status.DONE,
        )
        response = self.client.get(reverse("repairs:list"), {"status": Repair.Status.DONE})
        self.assertContains(response, "Другая заявка")
        self.assertNotContains(response, "Не набирает температуру")
        self.assertTrue(Repair.objects.filter(pk=other.pk).exists())

    def test_search_by_error_code(self):
        response = self.client.get(reverse("repairs:list"), {"q": "E4"})
        self.assertContains(response, "Не набирает температуру")

    def test_search_no_match(self):
        response = self.client.get(reverse("repairs:list"), {"q": "не найдётся"})
        self.assertNotContains(response, "Не набирает температуру")


class RepairExportViewsTests(RepairViewsTestCase):
    def setUp(self):
        super().setUp()
        self.repair = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Не набирает температуру",
            error_code="E4",
        )

    def test_csv_requires_login(self):
        self.client.logout()
        response = self.client.get(reverse("repairs:export_csv"))
        self.assertEqual(response.status_code, 302)

    def test_xlsx_requires_login(self):
        self.client.logout()
        response = self.client.get(reverse("repairs:export_xlsx"))
        self.assertEqual(response.status_code, 302)

    def test_csv_export_contains_visible_repair(self):
        response = self.client.get(reverse("repairs:export_csv"))
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"E4", response.content)
        self.assertIn("Не набирает температуру".encode(), response.content)

    def test_csv_export_respects_status_filter(self):
        other = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Другая заявка",
            status=Repair.Status.DONE,
        )
        response = self.client.get(
            reverse("repairs:export_csv"), {"status": Repair.Status.DONE}
        )
        content = response.content.decode("utf-8-sig")
        self.assertIn("Другая заявка", content)
        self.assertNotIn("Не набирает температуру", content)
        self.assertTrue(Repair.objects.filter(pk=other.pk).exists())

    def test_csv_export_excludes_archived_by_default(self):
        self.repair.is_archived = True
        self.repair.save(update_fields=["is_archived"])
        response = self.client.get(reverse("repairs:export_csv"))
        content = response.content.decode("utf-8-sig")
        self.assertNotIn("Не набирает температуру", content)

    def test_csv_export_includes_archived_when_requested(self):
        self.repair.is_archived = True
        self.repair.save(update_fields=["is_archived"])
        response = self.client.get(reverse("repairs:export_csv"), {"archived": "1"})
        content = response.content.decode("utf-8-sig")
        self.assertIn("Не набирает температуру", content)

    def test_xlsx_export_returns_valid_workbook(self):
        from io import BytesIO

        from openpyxl import load_workbook

        response = self.client.get(reverse("repairs:export_xlsx"))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 2)


class RepairDetailViewTests(RepairViewsTestCase):
    def test_shows_repair_info(self):
        repair = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Не набирает температуру",
        )
        response = self.client.get(reverse("repairs:detail", args=[repair.pk]))
        self.assertContains(response, "Не набирает температуру")

    def test_offers_call_and_max_buttons_for_client_with_phone(self):
        repair = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Не набирает температуру",
        )
        response = self.client.get(reverse("repairs:detail", args=[repair.pk]))
        self.assertContains(response, 'href="tel:+79990001122"')
        self.assertContains(response, "write-max-btn")

    def test_offers_no_phone_redirect_for_client_without_phone(self):
        client_without_phone = Client.objects.create(name="Кафе Без Номера")
        equipment = Equipment.objects.create(
            client=client_without_phone,
            equipment_type=self.equipment_type,
            brand=self.brand,
            model="Other",
        )
        repair = Repair.objects.create(
            equipment=equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Не набирает температуру",
        )
        response = self.client.get(reverse("repairs:detail", args=[repair.pk]))
        self.assertContains(response, "no-phone-btn")
        self.assertContains(
            response,
            f'data-edit-url="{reverse("clients:update", args=[client_without_phone.pk])}"',
        )

    def test_hides_invoice_section_before_completion(self):
        repair = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Не набирает температуру",
            status=Repair.Status.IN_PROGRESS,
        )
        response = self.client.get(reverse("repairs:detail", args=[repair.pk]))
        self.assertNotContains(response, "Счёт выставлен")
        self.assertNotContains(response, "Счёт не выставлен")

    def test_shows_invoice_status_after_completion(self):
        repair = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Не набирает температуру",
            status=Repair.Status.DONE,
        )
        response = self.client.get(reverse("repairs:detail", args=[repair.pk]))
        self.assertContains(response, "Счёт не выставлен")

        repair.invoice_issued = True
        repair.save(update_fields=["invoice_issued"])
        response = self.client.get(reverse("repairs:detail", args=[repair.pk]))
        self.assertContains(response, "Счёт выставлен")


class RepairNalogButtonTests(RepairViewsTestCase):
    def test_list_offers_nalog_button(self):
        response = self.client.get(reverse("repairs:list"))
        self.assertContains(response, "nalog-btn")

    def test_detail_nalog_button_carries_the_clients_inn(self):
        self.client_obj.inn = "770708389431"
        self.client_obj.save(update_fields=["inn"])
        repair = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Не набирает температуру",
        )
        response = self.client.get(reverse("repairs:detail", args=[repair.pk]))
        self.assertContains(response, 'data-inn="770708389431"')


class RepairInvoiceToggleViewTests(RepairViewsTestCase):
    def setUp(self):
        super().setUp()
        self.repair = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Не набирает температуру",
            status=Repair.Status.DONE,
        )

    def test_requires_login(self):
        self.client.logout()
        response = self.client.post(reverse("repairs:invoice_toggle", args=[self.repair.pk]))
        self.assertEqual(response.status_code, 302)

    def test_toggles_invoice_issued_on(self):
        response = self.client.post(reverse("repairs:invoice_toggle", args=[self.repair.pk]))
        self.assertRedirects(response, reverse("repairs:detail", args=[self.repair.pk]))
        self.repair.refresh_from_db()
        self.assertTrue(self.repair.invoice_issued)

    def test_toggles_invoice_issued_off_again(self):
        self.repair.invoice_issued = True
        self.repair.save(update_fields=["invoice_issued"])
        self.client.post(reverse("repairs:invoice_toggle", args=[self.repair.pk]))
        self.repair.refresh_from_db()
        self.assertFalse(self.repair.invoice_issued)


class RepairUpdateViewTests(RepairViewsTestCase):
    def test_updates_repair(self):
        repair = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Не набирает температуру",
        )
        response = self.client.post(
            reverse("repairs:update", args=[repair.pk]),
            {
                "master": self.user.pk,
                "status": Repair.Status.DIAGNOSTICS,
                "reported_at": repair.reported_at.isoformat(),
                "symptom": "Не набирает температуру",
                "error_code": "",
                "diagnosis": "Неисправен ТЭН",
                "work_done": "",
                "fault": "",
                "labor_cost": "0",
                "parts_cost": "0",
            },
        )
        self.assertRedirects(response, reverse("repairs:detail", args=[repair.pk]))
        repair.refresh_from_db()
        self.assertEqual(repair.status, Repair.Status.DIAGNOSTICS)
        self.assertEqual(repair.diagnosis, "Неисправен ТЭН")

    def test_error_code_field_wired_for_fault_suggestions(self):
        repair = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Не набирает температуру",
        )
        response = self.client.get(reverse("repairs:update", args=[repair.pk]))
        attrs = response.context["form"].fields["error_code"].widget.attrs
        self.assertEqual(attrs["hx-get"], reverse("knowledge:fault_suggestions"))
        self.assertIn(str(self.equipment_type.pk), attrs["hx-vals"])
        self.assertContains(response, 'id="fault-suggestions"')


class RepairStatusUpdateViewTests(RepairViewsTestCase):
    def setUp(self):
        super().setUp()
        self.repair = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate() - timedelta(days=1),
            symptom="Не набирает температуру",
        )

    def test_updates_status(self):
        response = self.client.post(
            reverse("repairs:status", args=[self.repair.pk]),
            {"status": Repair.Status.IN_PROGRESS},
        )
        self.assertRedirects(response, reverse("repairs:detail", args=[self.repair.pk]))
        self.repair.refresh_from_db()
        self.assertEqual(self.repair.status, Repair.Status.IN_PROGRESS)

    def test_marking_done_sets_completed_at(self):
        self.client.post(
            reverse("repairs:status", args=[self.repair.pk]),
            {"status": Repair.Status.DONE},
        )
        self.repair.refresh_from_db()
        self.assertEqual(self.repair.completed_at, timezone.localdate())


class RepairArchiveViewTests(RepairViewsTestCase):
    def setUp(self):
        super().setUp()
        self.repair = Repair.objects.create(
            equipment=self.equipment,
            master=self.user,
            reported_at=timezone.localdate(),
            symptom="Не набирает температуру",
        )

    def test_archive_confirmation_page_does_not_change_state(self):
        response = self.client.get(reverse("repairs:archive", args=[self.repair.pk]))
        self.assertEqual(response.status_code, 200)
        self.repair.refresh_from_db()
        self.assertFalse(self.repair.is_archived)

    def test_archive_post_soft_deletes(self):
        response = self.client.post(reverse("repairs:archive", args=[self.repair.pk]))
        self.assertRedirects(response, reverse("repairs:detail", args=[self.repair.pk]))
        self.repair.refresh_from_db()
        self.assertTrue(self.repair.is_archived)

    def test_restore_clears_archived_flag(self):
        self.repair.is_archived = True
        self.repair.save(update_fields=["is_archived"])
        response = self.client.post(reverse("repairs:restore", args=[self.repair.pk]))
        self.assertRedirects(response, reverse("repairs:detail", args=[self.repair.pk]))
        self.repair.refresh_from_db()
        self.assertFalse(self.repair.is_archived)
