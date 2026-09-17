import csv
import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from apps.clients.models import Client
from apps.equipment.models import Brand, Equipment, EquipmentType
from apps.repairs.export import EXPORT_HEADERS, export_repairs_csv, export_repairs_xlsx
from apps.repairs.models import Repair

User = get_user_model()


class RepairExportTestCase(TestCase):
    def setUp(self):
        self.master = User.objects.create_user(username="master", password="pass12345")
        client = Client.objects.create(name="Ресторан Восток")
        equipment_type = EquipmentType.objects.create(
            name="Пароконвектомат", category=EquipmentType.Category.KITCHEN
        )
        brand = Brand.objects.create(name="Rational")
        self.equipment = Equipment.objects.create(
            client=client,
            equipment_type=equipment_type,
            brand=brand,
            model="SCC 61",
            serial_number="SN-001",
        )
        self.repair = Repair.objects.create(
            equipment=self.equipment,
            master=self.master,
            reported_at=timezone.localdate(),
            symptom="Не набирает температуру",
            error_code="E4",
            labor_cost=Decimal("1500.00"),
            parts_cost=Decimal("750.50"),
        )


class ExportRepairsCSVTests(RepairExportTestCase):
    def test_content_type_and_filename(self):
        response = export_repairs_csv(Repair.objects.all())
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("attachment; filename=", response["Content-Disposition"])
        self.assertTrue(response["Content-Disposition"].endswith('.csv"'))

    def test_starts_with_utf8_bom(self):
        response = export_repairs_csv(Repair.objects.all())
        self.assertTrue(response.content.decode("utf-8").startswith("﻿"))

    def test_header_row_matches(self):
        response = export_repairs_csv(Repair.objects.all())
        content = response.content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content))
        header = next(reader)
        self.assertEqual(header, EXPORT_HEADERS)

    def test_data_row_has_expected_values(self):
        response = export_repairs_csv(Repair.objects.all())
        content = response.content.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(content))
        next(reader)
        row = next(reader)
        self.assertEqual(row[0], str(self.repair.pk))
        self.assertEqual(row[3], "Ресторан Восток")
        self.assertEqual(row[6], "E4")
        self.assertEqual(row[7], "Не набирает температуру")
        self.assertEqual(row[13], "2250.50")

    def test_empty_queryset_only_has_header(self):
        response = export_repairs_csv(Repair.objects.none())
        content = response.content.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(content)))
        self.assertEqual(len(rows), 1)


class ExportRepairsXLSXTests(RepairExportTestCase):
    def test_content_type_and_filename(self):
        response = export_repairs_xlsx(Repair.objects.all())
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertTrue(response["Content-Disposition"].endswith('.xlsx"'))

    def test_workbook_has_header_and_data_row(self):
        response = export_repairs_xlsx(Repair.objects.all())
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active

        header = [cell.value for cell in sheet[1]]
        self.assertEqual(header, EXPORT_HEADERS)

        data_row = [cell.value for cell in sheet[2]]
        self.assertEqual(data_row[0], self.repair.pk)
        self.assertEqual(data_row[3], "Ресторан Восток")
        self.assertEqual(data_row[6], "E4")

    def test_header_row_is_bold(self):
        response = export_repairs_xlsx(Repair.objects.all())
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        self.assertTrue(sheet["A1"].font.bold)

    def test_sheet_row_count_matches_queryset(self):
        Repair.objects.create(
            equipment=self.equipment,
            master=self.master,
            reported_at=timezone.localdate(),
            symptom="Другая заявка",
        )
        response = export_repairs_xlsx(Repair.objects.all())
        workbook = load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        # header + 2 data rows
        self.assertEqual(sheet.max_row, 3)
