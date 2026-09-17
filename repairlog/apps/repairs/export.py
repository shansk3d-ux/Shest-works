import csv
import io

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

EXPORT_HEADERS = [
    "№",
    "Дата обращения",
    "Статус",
    "Клиент",
    "Оборудование",
    "Серийный номер",
    "Код ошибки",
    "Симптом",
    "Диагноз",
    "Выполненные работы",
    "Мастер",
    "Стоимость работ",
    "Стоимость запчастей",
    "Итого",
    "Дата завершения",
    "Гарантия до",
]


def repair_export_row(repair):
    return [
        repair.pk,
        repair.reported_at.isoformat() if repair.reported_at else "",
        repair.get_status_display(),
        repair.equipment.client.name,
        str(repair.equipment),
        repair.equipment.serial_number,
        repair.error_code,
        repair.symptom,
        repair.diagnosis,
        repair.work_done,
        repair.master.get_username(),
        repair.labor_cost,
        repair.parts_cost,
        repair.total_cost,
        repair.completed_at.isoformat() if repair.completed_at else "",
        repair.warranty_until.isoformat() if repair.warranty_until else "",
    ]


def _export_filename(extension):
    timestamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    return f"repairs-{timestamp}.{extension}"


def export_repairs_csv(queryset):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(EXPORT_HEADERS)
    for repair in queryset:
        writer.writerow(repair_export_row(repair))

    # Prepend a UTF-8 BOM so Excel (which otherwise assumes the system codepage)
    # displays Cyrillic text correctly instead of mangling it.
    response = HttpResponse("﻿" + buffer.getvalue(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{_export_filename("csv")}"'
    return response


def export_repairs_xlsx(queryset):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заявки"

    sheet.append(EXPORT_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for repair in queryset:
        sheet.append(repair_export_row(repair))

    for index, header in enumerate(EXPORT_HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(12, min(40, len(header) + 6))

    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{_export_filename("xlsx")}"'
    return response
